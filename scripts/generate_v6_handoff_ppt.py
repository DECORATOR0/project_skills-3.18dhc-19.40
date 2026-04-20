from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = 13.333
SLIDE_H = 7.5

FONT = "Noto Sans CJK SC"

BG = RGBColor(246, 243, 236)
TEXT = RGBColor(33, 38, 44)
MUTED = RGBColor(90, 99, 112)
CARD = RGBColor(255, 252, 247)
LINE = RGBColor(220, 211, 200)
ACCENT = RGBColor(191, 105, 55)
ACCENT_2 = RGBColor(45, 94, 131)
ACCENT_3 = RGBColor(58, 127, 118)
WARN = RGBColor(130, 45, 32)
GOOD = RGBColor(43, 112, 82)


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _style_run(run, font_size, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text="",
    font_size=20,
    bold=False,
    color=TEXT,
    align=PP_ALIGN.LEFT,
    valign=MSO_VERTICAL_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    if p.runs:
        _style_run(p.runs[0], font_size, bold=bold, color=color)
    return box


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    bullets,
    font_size=16,
    color=TEXT,
    prefix="• ",
    space_after=5,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{prefix}{item}"
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = 1.12
        if p.runs:
            _style_run(p.runs[0], font_size, color=color)
    return box


def add_card(
    slide,
    left,
    top,
    width,
    height,
    title,
    lines,
    accent=ACCENT,
    title_size=19,
    body_size=16,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1.2)

    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.12),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    body_top = top + 0.56
    body_height = height - 0.7
    if title:
        add_textbox(slide, left + 0.18, top + 0.18, width - 0.36, 0.34, title, title_size, True, TEXT)
    else:
        body_top = top + 0.24
        body_height = height - 0.34

    add_bullets(
        slide,
        left + 0.18,
        body_top,
        width - 0.36,
        body_height,
        lines,
        font_size=body_size,
        prefix="· ",
    )


def add_title(slide, title, subtitle=None):
    set_bg(slide)
    add_textbox(slide, 0.72, 0.48, 11.9, 0.58, title, 27, True, TEXT)
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.74),
        Inches(1.18),
        Inches(2.5),
        Inches(0.08),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if subtitle:
        add_textbox(slide, 0.76, 1.42, 11.7, 0.52, subtitle, 15, False, MUTED)


def add_footer_note(slide, text):
    add_textbox(slide, 0.76, 6.92, 11.9, 0.28, text, 10, False, MUTED, align=PP_ALIGN.RIGHT)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(SLIDE_W),
        Inches(1.08),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT_2
    band.line.fill.background()

    add_textbox(slide, 0.72, 1.4, 11.8, 0.8, "V6 后续复盘与下一步 benchmark 计划", 29, True, TEXT)
    add_textbox(
        slide,
        0.76,
        2.18,
        11.2,
        0.7,
        "面向 V6 开发者的 handoff 版本：先接上现状，再对齐下一轮主战场",
        16,
        False,
        MUTED,
    )

    add_card(
        slide,
        0.78,
        4.05,
        3.72,
        1.72,
        "汇报目标",
        ["让对方快速接上 4/13 以来的变化", "用最少细节说明为什么需要换主 benchmark"],
        ACCENT,
        body_size=15,
    )
    add_card(
        slide,
        4.81,
        4.05,
        3.72,
        1.72,
        "当前判断",
        ["主要瓶颈已落到 gold / data / runtime contract", "skill 侧有效修正已经做过多轮"],
        ACCENT_3,
        body_size=15,
    )
    add_card(
        slide,
        8.84,
        4.05,
        3.72,
        1.72,
        "这次重点",
        ["先复盘清楚发生了什么", "再明确接下来在哪个 benchmark 上继续做"],
        ACCENT_2,
        body_size=15,
    )
    add_footer_note(slide, "内部 handoff 用，细节版文档另附")


def add_timeline_slide_a(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "从 V6 到当前：演化线索（一）", "前半段主要在把 V6 prompt 接入本地框架，并把执行约束收紧")

    add_card(
        slide,
        0.82,
        2.0,
        5.85,
        3.2,
        "4/13：接上 V6",
        [
            "接入对方 V6 prompt，到本地 batch single-skill 框架。",
            "固定 phase-based single skill 执行骨架。",
            "早期结果 14/30，但其中混有不少脆弱成功。",
        ],
        ACCENT,
    )
    add_card(
        slide,
        6.75,
        2.0,
        5.75,
        3.2,
        "4/14 白天：先把执行守住",
        [
            "strict-state 接管 phase 跳转。",
            "禁止越 phase 提前作答。",
            "执行更规整，同时暴露出更多假阳性。",
        ],
        ACCENT_2,
    )
    add_card(
        slide,
        0.82,
        5.55,
        11.68,
        0.92,
        "这一阶段的重点",
        ["先把“能跑”变成“按约束跑”，让后面的分数更接近真实能力。"],
        GOOD,
        body_size=16,
    )
    add_footer_note(slide, "时间线 1/2")


def add_timeline_slide_b(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "从 V6 到当前：演化线索（二）", "后半段开始同时修 skill 修改链、runtime 契约和数据对齐")

    add_card(
        slide,
        0.82,
        2.0,
        5.85,
        3.2,
        "4/14 晚上：skill 修改链补齐",
        [
            "修 batch 工具契约。",
            "修 actor / critic 输入和修改策略。",
            "skill 已能更定向地修局部问题。",
        ],
        ACCENT_3,
    )
    add_card(
        slide,
        6.75,
        2.0,
        5.75,
        3.2,
        "4/15：做对照、做数据修补",
        [
            "修 semantic success 与产物路径。",
            "做 dataalignv1，并补 gold overrides。",
            "加 strict gold replay 与 GPT-5.2 裸跑对照。",
        ],
        WARN,
    )
    add_card(
        slide,
        0.82,
        5.55,
        11.68,
        0.92,
        "总体趋势",
        ["约束逐步收紧后，虚高成功被挤掉；真实上限开始稳定在 17/18 左右。"],
        ACCENT,
        body_size=16,
    )
    add_footer_note(slide, "时间线 2/2")


def add_delta_slide_a(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "这几天实际做过的修正（一）", "这一页只保留与原始 V6 差异最大的两块")

    add_card(
        slide,
        0.82,
        2.0,
        5.85,
        3.3,
        "Skill / Executor",
        [
            "从 V6 prompt 落到 phase-based single skill。",
            "strict-state 显式接管 SEARCH / VERIFY / CONCLUDE。",
            "只有 CONCLUDE 允许输出最终答案。",
        ],
        ACCENT,
    )
    add_card(
        slide,
        6.75,
        2.0,
        5.75,
        3.3,
        "Actor / Critic",
        [
            "critic 读取完整 skill、题干与选项。",
            "gold args 做减重，减少无关噪声。",
            "actor 支持 ADD / REPLACE / DELETE，修改更细粒度。",
        ],
        ACCENT_2,
    )
    add_card(
        slide,
        0.82,
        5.58,
        11.68,
        0.9,
        "这一页想传达的重点",
        ["这两块决定了“skill 怎么改”和“executor 怎么守约束”，当前状态已经明显偏离原始 V6。"],
        ACCENT,
        body_size=16,
    )
    add_footer_note(slide, "版本差异 1/2")


def add_delta_slide_b(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "这几天实际做过的修正（二）", "后续看分数时，必须把 runtime 和 data 的变化一起考虑进去")

    add_card(
        slide,
        0.82,
        2.0,
        5.85,
        3.3,
        "Tools / Runtime",
        [
            "补 batch 双态接口与工具契约。",
            "修 semantic success 判定。",
            "补 SAM2 / ChangeOS 等产物路径映射。",
        ],
        ACCENT_3,
    )
    add_card(
        slide,
        6.75,
        2.0,
        5.75,
        3.3,
        "Data",
        [
            "系统扫 248 题。",
            "修多处 gold overrides。",
            "落地 dataalignv1，对齐题面 / 数据 / gold。",
        ],
        WARN,
    )
    add_card(
        slide,
        0.82,
        5.58,
        11.68,
        0.9,
        "这一页想传达的重点",
        ["后面再看分数，已经不能直接理解成“同一个 V6 在重复跑”。"],
        ACCENT,
        body_size=16,
    )
    add_footer_note(slide, "版本差异 2/2")


def add_issue_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "当前问题已经收敛成四层", "这页只讲主问题，细枝末节全部压缩")

    add_card(
        slide,
        0.82,
        1.95,
        5.85,
        2.05,
        "1. Gold Answer 脏",
        [
            "248 题里已有 24 题高置信异常。",
            "做对也可能被判错。",
            "reward 会混入脏监督。",
        ],
        WARN,
        body_size=15,
    )
    add_card(
        slide,
        6.75,
        1.95,
        5.75,
        2.05,
        "2. Gold Trajectory 脏",
        [
            "strict replay 只有 1/30 完整通过。",
            "大量路径失效、语义漂移、旧产物依赖。",
            "部分题已经超出单点 patch 范围。",
        ],
        ACCENT_2,
        body_size=15,
    )
    add_card(
        slide,
        0.82,
        4.25,
        5.85,
        2.05,
        "3. Data / Prompt 错配",
        [
            "248 题里有 26 题题面和真实数据不一致。",
            "题面、目录、gold 往往不是同一套东西。",
            "executor 很容易被带偏。",
        ],
        ACCENT_3,
        body_size=15,
    )
    add_card(
        slide,
        6.75,
        4.25,
        5.75,
        2.05,
        "4. Tool Contract 脆弱",
        [
            "63/248 题命中 batch 契约冲突。",
            "历史成功里混有假成功。",
            "修完后暴露出更多真实失败。",
        ],
        ACCENT,
        body_size=15,
    )
    add_footer_note(slide, "四类问题共同吞噬收益")


def add_evidence_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "为什么这个判断已经比较稳定", "四组数字已经足够把现状钉住")

    add_card(
        slide,
        0.82,
        1.95,
        5.85,
        2.02,
        "早期迭代：14/30",
        ["其中只有约 7 题能算稳成功，另一半更多是脆弱命中。"],
        ACCENT,
        body_size=16,
    )
    add_card(
        slide,
        6.75,
        1.95,
        5.75,
        2.02,
        "约束收紧后：17-18/30",
        ["状态机、toolfix、prompt 修完后，长期在这个区间波动。"],
        ACCENT_2,
        body_size=16,
    )
    add_card(
        slide,
        0.82,
        4.22,
        5.85,
        2.02,
        "strict gold replay：1/30",
        ["直接按 gold 轨迹回放几乎全灭，问题不只落在 executor 一侧。"],
        WARN,
        body_size=16,
    )
    add_card(
        slide,
        6.75,
        4.22,
        5.75,
        2.02,
        "GPT-5.2 no-skill：14/30",
        ["更强模型裸跑上限也不高，说明 benchmark 本身限制很重。"],
        ACCENT_3,
        body_size=16,
    )
    add_footer_note(slide, "结论：继续深挖 EO，主要成本会落到脏数据与脏 gold 清洗")


def add_judgement_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "对 EO Bench 的当前定位", "需要把“诊断集”与“主战场”明确区分开")

    add_card(
        slide,
        0.82,
        2.0,
        5.85,
        3.2,
        "EO Bench 仍然有价值",
        [
            "适合做诊断集与回归集。",
            "适合继续盯 runtime contract、路径与 success 判定。",
            "适合固化 batch30 + strict replay + 代表性脏题子集。",
        ],
        GOOD,
    )
    add_card(
        slide,
        6.75,
        2.0,
        5.75,
        3.2,
        "EO Bench 不适合继续做主战场",
        [
            "题型偏工具顺序和稳健性，skill 信号不够纯。",
            "benchmark 噪声会遮住真实收益。",
            "继续深挖，时间会主要花在清洗脏 benchmark。",
        ],
        WARN,
    )
    add_card(
        slide,
        0.82,
        5.55,
        11.68,
        0.92,
        "更合理的定位",
        ["EO Bench 保留为诊断包；主实验尽快切到新的 benchmark。"],
        ACCENT,
        body_size=16,
    )
    add_footer_note(slide, "诊断包继续保留，主实验切走")


def add_benchmark_overview_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "新的主 benchmark 候选：先看 shortlist", "按“能否体现 skill 架构独特性 + 迁移成本”做压缩")

    add_card(
        slide,
        0.82,
        1.95,
        5.85,
        2.1,
        "主选 1：BrowseComp",
        ["浏览检索 + 多跳约束核验。", "直接测 decomposition、fetch、停止条件。"],
        ACCENT,
        body_size=16,
    )
    add_card(
        slide,
        6.75,
        1.95,
        5.75,
        2.1,
        "主选 2：SealQA",
        ["noisy retrieval QA。", "直接测消歧、冲突证据处理、拒答。"],
        ACCENT_2,
        body_size=16,
    )
    add_card(
        slide,
        0.82,
        4.3,
        5.85,
        2.1,
        "备选：OfficeQA Pro",
        ["长文档 + 表格 + 数值推理。", "适合测 fetch、结构化抽取和轻量计算。"],
        ACCENT_3,
        body_size=16,
    )
    add_card(
        slide,
        6.75,
        4.3,
        5.75,
        2.1,
        "验证线：SkillsBench",
        ["显式 skill + verifier。", "适合补做“skill 是否真有收益”的对照。"],
        GOOD,
        body_size=16,
    )
    add_footer_note(slide, "GAIA 放第二阶段；SWE-bench 暂不承担当前主结论")


def add_browsecomp_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "主选 1：BrowseComp", "如果目标是最快看到 skill 信号，这个 benchmark 很适合先落")

    add_card(
        slide,
        0.82,
        1.95,
        4.2,
        3.6,
        "为什么首推",
        [
            "任务天然依赖多步检索与约束拆解。",
            "最终答案通常很短，误差主要集中在搜索与验证阶段。",
            "fetch、渐进式披露、状态跳转都能发挥作用。",
        ],
        ACCENT,
    )
    add_card(
        slide,
        5.22,
        1.95,
        7.28,
        3.6,
        "每题怎么跑",
        [
            "先拆题：列出实体、时间、关系等关键约束。",
            "先粗搜，再用更具体的查询补齐缺失约束。",
            "在多个页面间交叉核验，直到只剩一个候选答案。",
            "用短答案收尾，并给出简短证据摘要或置信判断。",
        ],
        ACCENT_2,
        body_size=15,
    )
    add_card(
        slide,
        0.82,
        5.6,
        11.68,
        1.12,
        "代表性示意题",
        [
            "示意题：已知 1998 世界杯决赛最后进球者在两年后与饰演 Victoria 的演员结婚，问该演员 1995 年主演的电影名；关键看能否持续拆约束并回查关系与时间线。",
        ],
        ACCENT,
        body_size=14,
    )


def add_sealqa_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "主选 2：SealQA", "如果想更直接测“脏检索环境下的 skill 价值”，SealQA 很合适")

    add_card(
        slide,
        0.82,
        1.95,
        4.2,
        3.6,
        "为什么首推",
        [
            "搜索结果天然有噪声、冲突和歧义。",
            "很适合测 strict-state 下的“继续搜 / 改查询 / 拒答”。",
            "能看出 skill 是否真能帮助模型处理脏检索环境。",
        ],
        ACCENT_2,
    )
    add_card(
        slide,
        5.22,
        1.95,
        7.28,
        3.6,
        "每题怎么跑",
        [
            "先直接搜原题，观察冲突点落在哪个实体、年份或地点条件。",
            "改写查询做消歧：补年份、组织名、地理约束、官方来源。",
            "对比多源证据，判断该回答、该拒答，还是该指出前提有误。",
            "只在证据闭环后收尾。",
        ],
        ACCENT_3,
        body_size=15,
    )
    add_card(
        slide,
        0.82,
        5.6,
        11.68,
        1.12,
        "代表性示意题",
        [
            "示意题：问“2024 年获某奖的 John Williams 出生于哪里”，搜索结果混入同名作曲家、摄影师和旧年份页面；关键看能否先消歧，再围绕时间与身份做定向检索。",
        ],
        ACCENT_2,
        body_size=14,
    )


def add_officeqa_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "备选：OfficeQA Pro", "如果想测 fetch、文档定位与轻量计算链条，这条线更合适")

    add_card(
        slide,
        0.82,
        1.95,
        4.2,
        3.6,
        "为什么可选",
        [
            "文档长、表格密、数值链条清晰。",
            "很适合测 fetch、页面定位、结构化抽取和轻量计算。",
            "结果可验证，噪声相对 EO Bench 小很多。",
        ],
        ACCENT_3,
    )
    add_card(
        slide,
        5.22,
        1.95,
        7.28,
        3.6,
        "每题怎么跑",
        [
            "先根据题目定位年份、主题和对应 Treasury Bulletin。",
            "在 PDF 或转换文本里找到相关章节、表格、列头和单位。",
            "抽出关键数值后做同比 / 环比 / 比率等简单计算。",
            "输出短答案，必要时附页码或文件名。",
        ],
        ACCENT,
        body_size=15,
    )
    add_card(
        slide,
        0.82,
        5.6,
        11.68,
        1.12,
        "代表性示意题",
        [
            "示意题：在 1941 年 1 月与 1946 年 1 月的 Treasury Bulletin 中找到同一项国债指标，计算百分比变化；关键看能否找对页、读对表、算对值。",
        ],
        ACCENT_3,
        body_size=14,
    )


def add_other_candidates_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "其他候选怎么放置", "SkillsBench、GAIA、SWE-bench 都有价值，但位置需要分清")

    add_card(
        slide,
        0.82,
        2.0,
        3.72,
        3.45,
        "SkillsBench",
        [
            "价值：直接测显式 skill / verifier 对结果的提升。",
            "问题：环境和评测体系偏重，首轮迁移成本高。",
            "建议：主 benchmark 跑稳后，再并行做验证线。",
        ],
        GOOD,
        body_size=15,
    )
    add_card(
        slide,
        4.81,
        2.0,
        3.72,
        3.45,
        "GAIA",
        [
            "价值：通用 agent 泛化，附件 + 网页 + 多步工具都能测。",
            "问题：题型异质性高，首轮更容易先测出通用 scaffold。",
            "建议：放第二阶段做泛化压力测试。",
        ],
        ACCENT_2,
        body_size=15,
    )
    add_card(
        slide,
        8.8,
        2.0,
        3.72,
        3.45,
        "SWE-bench",
        [
            "价值：社区认知高，对外解释强。",
            "问题：coding scaffold、基础模型能力、repo 工程因素权重都很高。",
            "建议：另开 coding 支线，暂不承担当前主结论。",
        ],
        WARN,
        body_size=15,
    )
    add_footer_note(slide, "首轮主线仍建议放在 BrowseComp / SealQA")


def add_plan_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "接下来怎么推进", "把这次复盘直接转成可执行动作")

    add_card(
        slide,
        0.82,
        1.95,
        3.72,
        3.45,
        "Step 1：收口当前线",
        [
            "把 EO Bench 固化成诊断包。",
            "固定 batch30、strict replay、代表性脏题。",
            "后续任何修改都先过这组回归。",
        ],
        ACCENT,
        body_size=15,
    )
    add_card(
        slide,
        4.81,
        1.95,
        3.72,
        3.45,
        "Step 2：定主线",
        [
            "先在 BrowseComp 和 SealQA 上各跑一个 20-30 题最小闭环。",
            "对比 no-skill / V6 / 当前 strict-state 版本。",
            "优先选择 skill 信号更清楚的一条线继续做。",
        ],
        ACCENT_2,
        body_size=15,
    )
    add_card(
        slide,
        8.8,
        1.95,
        3.72,
        3.45,
        "Step 3：第二轮扩展",
        [
            "围绕 fetch、渐进式披露、状态设计继续做增益。",
            "GAIA 只做泛化抽查。",
            "SkillsBench 视资源并行，SWE-bench 另开线。",
        ],
        ACCENT_3,
        body_size=15,
    )
    add_card(
        slide,
        0.82,
        5.72,
        11.68,
        0.9,
        "希望这次 handoff 帮对方接上的状态",
        ["现在最需要一起确认的是：主 benchmark 先选 BrowseComp 还是 SealQA，以及哪些 V6 假设必须保留。"],
        ACCENT,
        body_size=15,
    )
    add_footer_note(slide, "建议先定主 benchmark，再继续改 skill")


def build_deck(output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    add_title_slide(prs)
    add_timeline_slide_a(prs)
    add_timeline_slide_b(prs)
    add_delta_slide_a(prs)
    add_delta_slide_b(prs)
    add_issue_slide(prs)
    add_evidence_slide(prs)
    add_judgement_slide(prs)
    add_benchmark_overview_slide(prs)
    add_browsecomp_slide(prs)
    add_sealqa_slide(prs)
    add_officeqa_slide(prs)
    add_other_candidates_slide(prs)
    add_plan_slide(prs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main():
    parser = argparse.ArgumentParser(description="Generate the V6 handoff PPT.")
    parser.add_argument("output", nargs="?", help="Output .pptx path")
    args = parser.parse_args()

    default_output = Path(
        "/data/xsy/project_skills-3.18dhc-19.40/实验设计与迭代/26.4.16_2258_V6后续复盘与benchmark计划_handoff版_v2.pptx"
    )
    output = Path(args.output) if args.output else default_output
    build_deck(output)
    print(output)


if __name__ == "__main__":
    main()
